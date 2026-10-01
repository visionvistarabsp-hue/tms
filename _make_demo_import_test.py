import os, sys
sys.path.insert(0, r"C:\Users\shyam\Music\tms2-main\tms2-main")
os.chdir(r"C:\Users\shyam\Music\tms2-main\tms2-main")

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime

wb = Workbook()
ws = wb.active
ws.title = "TDS"

headers = ["Date", "Lorry Number", "Account Name", "Mines Qty", "Plant Qty", "Rate", "Cash"]
ws.append(headers)
for c in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=c)
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="1F4E79")
    cell.alignment = Alignment(horizontal="center")

rows = [
    (datetime(2026, 8, 1), "CG31B9931", "AMIT SONI", 42.6708, 36.2210, 420, 5600),
    (datetime(2026, 8, 2), "CG31B9931", "AMIT SONI", 41.9087, 35.5050, 415, 5500),
    (datetime(2026, 8, 3), "MP09AB1234", "BALAJI TRANSPORT", 40.1542, 34.8800, 410, 5400),
    (datetime(2026, 8, 4), "CG31B9931", "RAMESH KUMAR", 39.7718, 33.9900, 405, 5300),
]
for r in rows:
    ws.append(r)

for r in range(2, len(rows) + 2):
    for c in range(1, len(headers) + 1):
        ws.cell(row=r, column=c).number_format = (
            "yyyy-mm-dd" if c == 1 else
            ("0.0000" if headers[c - 1] in ("Mines Qty", "Plant Qty", "Rate") else "General")
        )

wb.save("demo_import_test.xlsx")
print("saved:", os.path.abspath("demo_import_test.xlsx"))
print("dims:", ws.dimensions)
